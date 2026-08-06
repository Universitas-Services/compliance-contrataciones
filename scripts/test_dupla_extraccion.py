"""Pruebas rápidas sin LLM: registro de 14 agentes + extracción canónica."""

from __future__ import annotations

import tempfile
from pathlib import Path

from app.agents.document_normalizer import partir_en_chunks
from app.agents.extractor import procesar_archivo
from app.agents.modalities import get_agente, listar_agentes_registrados
from app.agents.modalities.prompts import ANALISTA_PROMPTS, JURIDICO_PROMPTS
from app.models.schemas import Modalidad, RolAgente, TipoContratacion


def test_14_agentes() -> None:
    regs = listar_agentes_registrados()
    assert len(regs) == 14, regs
    for m in Modalidad:
        assert (m.value, RolAgente.ANALISTA.value) in regs
        assert (m.value, RolAgente.JURIDICO.value) in regs
        a = get_agente(m, TipoContratacion.BIENES, RolAgente.ANALISTA)
        j = get_agente(m, TipoContratacion.OBRAS, RolAgente.JURIDICO)
        assert a.rol == RolAgente.ANALISTA
        assert j.rol == RolAgente.JURIDICO
        assert "ANALISTA" in a.system_prompt_experto().upper() or "Analista" in a.system_prompt_experto()
        assert "JURÍDICO" in j.system_prompt_experto() or "JURIDICO" in j.system_prompt_experto().upper()
        assert m in ANALISTA_PROMPTS and m in JURIDICO_PROMPTS
    print("OK 14 agentes + prompts")


def test_chunks() -> None:
    big = "## p.1\n" + ("palabra " * 4000) + "\n## p.2\n" + ("otra " * 4000)
    parts = partir_en_chunks(big, max_chars=5000)
    assert len(parts) >= 2
    assert sum(len(p) for p in parts) >= len(big) - 50
    print("OK chunking", len(parts), "bloques")


def test_docx_extract() -> None:
    from docx import Document

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "demo.docx"
        doc = Document()
        doc.add_heading("Contrato de prueba", level=1)
        doc.add_paragraph("Las partes acuerdan el objeto X bajo nomenclatura DEMO-001.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Cláusula"
        table.rows[0].cells[1].text = "Valor"
        table.rows[1].cells[0].text = "Plazo"
        table.rows[1].cells[1].text = "30 días"
        doc.save(path)
        out = procesar_archivo(str(path), "demo.docx")
        assert out["modo"] in {"texto", "mixto"}
        assert "Contrato" in out["contenido"] or "Contrato" in out.get("texto_extraido", "")
        assert out["canonico"]["metodo"] == "docx"
        print("OK docx extract", out["canonico"])


if __name__ == "__main__":
    test_14_agentes()
    test_chunks()
    test_docx_extract()
    print("Todas las pruebas locales OK")
