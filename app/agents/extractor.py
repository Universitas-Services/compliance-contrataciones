"""Extracción canónica: texto/tablas/OCR multipágina (visión solo selectiva)."""

from __future__ import annotations

import base64
import mimetypes
from pathlib import Path
from typing import Any, Literal

import pandas as pd
import pdfplumber
from docx import Document

from app.agents.document_normalizer import (
    BloqueDocumento,
    DocumentoCanonico,
    armar_texto_canonico,
    chunks_desde_bloques,
    ocr_archivo_imagen,
    ocr_pdf_todas_paginas,
    partir_en_chunks,
    rasterizar_pagina_pdf,
)

ModoContenido = Literal["texto", "imagen", "mixto"]
TipoArchivo = Literal["pdf", "docx", "xlsx", "imagen", "desconocido"]

_EXTENSIONES_IMAGEN = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_TEXTO_MINIMO = 40


def detectar_tipo_archivo(filename: str) -> TipoArchivo:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in {".xlsx", ".xls"}:
        return "xlsx"
    if ext in _EXTENSIONES_IMAGEN:
        return "imagen"
    return "desconocido"


def _tablas_pagina_tsv(page) -> str:
    try:
        tables = page.extract_tables() or []
    except Exception:  # noqa: BLE001
        return ""
    partes: list[str] = []
    for ti, table in enumerate(tables, start=1):
        filas = []
        for row in table:
            celdas = [(c or "").replace("\t", " ").replace("\n", " ").strip() for c in row]
            if any(celdas):
                filas.append("\t".join(celdas))
        if filas:
            partes.append(f"[tabla {ti}]\n" + "\n".join(filas))
    return "\n".join(partes)


def extraer_pdf_nativo(path: str) -> DocumentoCanonico | None:
    bloques: list[BloqueDocumento] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            texto = (page.extract_text() or "").strip()
            tablas = _tablas_pagina_tsv(page)
            cuerpo = texto
            if tablas:
                cuerpo = f"{texto}\n\n{tablas}".strip() if texto else tablas
            if cuerpo:
                bloques.append(
                    BloqueDocumento(id=f"b{i}", etiqueta=f"p.{i}", texto=cuerpo)
                )
    texto_join = "\n\n".join(b.texto for b in bloques).strip()
    if len(texto_join) < _TEXTO_MINIMO:
        return None
    meta = [
        "# meta",
        f"archivo_tipo: pdf",
        f"metodo: pdfplumber+tablas",
        f"paginas: {len(bloques)}",
    ]
    return DocumentoCanonico(
        texto=armar_texto_canonico(meta_lineas=meta, bloques=bloques),
        metodo="pdf_nativo",
        paginas=len(bloques),
        bloques=bloques,
    )


def extraer_pdf_ocr(path: str) -> DocumentoCanonico:
    bloques, adv, cobertura = ocr_pdf_todas_paginas(path)
    meta = [
        "# meta",
        "archivo_tipo: pdf",
        "metodo: rapidocr_multipagina",
        f"paginas: {len(bloques)}",
        f"cobertura_ocr: {cobertura:.2f}",
    ]
    imagenes: list[dict] = []
    # Visión selectiva: páginas con OCR débil (máx 3)
    debiles = [a for a in adv if "OCR débil" in a][:3]
    for a in debiles:
        # extrae número de página
        try:
            num = int(a.split("p.")[-1].strip()) - 1
            png = rasterizar_pagina_pdf(path, num, zoom=1.5)
            imagenes.append(
                {
                    "type": "image",
                    "source": {
                        "type": "base64",
                        "media_type": "image/png",
                        "data": base64.standard_b64encode(png).decode("ascii"),
                    },
                    "pagina": num + 1,
                }
            )
        except Exception:  # noqa: BLE001
            continue
    if debiles:
        adv.append(
            f"Visión selectiva adjunta para {len(imagenes)} página(s) con OCR débil"
        )
    return DocumentoCanonico(
        texto=armar_texto_canonico(meta_lineas=meta, bloques=bloques),
        metodo="pdf_ocr",
        paginas=len(bloques),
        advertencias=adv,
        cobertura_ocr=cobertura,
        bloques=bloques,
        imagenes_selectivas=imagenes,
    )


def extraer_docx(path: str) -> DocumentoCanonico:
    doc = Document(path)
    partes: list[str] = []
    for section in doc.sections:
        for hf_name, hf in (
            ("encabezado", section.header),
            ("pie", section.footer),
        ):
            textos = [p.text.strip() for p in hf.paragraphs if p.text.strip()]
            if textos:
                partes.append(f"[{hf_name}]\n" + "\n".join(textos))
    for parrafo in doc.paragraphs:
        texto = parrafo.text.strip()
        if texto:
            style = (parrafo.style.name or "") if parrafo.style else ""
            if style.lower().startswith("heading"):
                partes.append(f"## {texto}")
            else:
                partes.append(texto)
    for ti, tabla in enumerate(doc.tables, start=1):
        filas = []
        for fila in tabla.rows:
            celdas = [c.text.strip().replace("\t", " ") for c in fila.cells]
            if any(celdas):
                filas.append("\t".join(celdas))
        if filas:
            partes.append(f"[tabla {ti}]\n" + "\n".join(filas))
    bloque = BloqueDocumento(id="b1", etiqueta="docx", texto="\n\n".join(partes))
    meta = ["# meta", "archivo_tipo: docx", "metodo: python-docx", "paginas: 1"]
    return DocumentoCanonico(
        texto=armar_texto_canonico(meta_lineas=meta, bloques=[bloque]),
        metodo="docx",
        paginas=1,
        bloques=[bloque],
    )


def extraer_xlsx(path: str) -> DocumentoCanonico:
    hojas = pd.read_excel(path, sheet_name=None, dtype=str)
    bloques: list[BloqueDocumento] = []
    for idx, (nombre, df) in enumerate(hojas.items(), start=1):
        df = df.fillna("")
        # quitar filas/columnas vacías
        df = df.loc[:, (df.astype(str).apply(lambda s: s.str.strip() != "")).any()]
        df = df.loc[(df.astype(str).apply(lambda r: r.str.strip() != "")).any(axis=1)]
        tsv = df.to_csv(index=False, sep="\t")
        bloques.append(
            BloqueDocumento(id=f"h{idx}", etiqueta=f"hoja:{nombre}", texto=tsv.strip())
        )
    meta = [
        "# meta",
        "archivo_tipo: xlsx",
        "metodo: pandas_tsv",
        f"paginas: {len(bloques)}",
    ]
    return DocumentoCanonico(
        texto=armar_texto_canonico(meta_lineas=meta, bloques=bloques),
        metodo="xlsx",
        paginas=len(bloques),
        bloques=bloques,
    )


def preparar_imagen(path: str) -> dict:
    ruta = Path(path)
    media_type, _ = mimetypes.guess_type(ruta.name)
    if not media_type or not media_type.startswith("image/"):
        ext = ruta.suffix.lower()
        fallback = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".gif": "image/gif",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
            ".tif": "image/tiff",
            ".tiff": "image/tiff",
        }
        media_type = fallback.get(ext, "image/png")
    data = base64.standard_b64encode(ruta.read_bytes()).decode("ascii")
    return {
        "type": "image",
        "source": {
            "type": "base64",
            "media_type": media_type,
            "data": data,
        },
    }


def procesar_archivo(path: str, filename: str) -> dict[str, Any]:
    """
    Normaliza el archivo a representación canónica densa.

    Retorno:
      modo: texto|mixto|imagen
      contenido: str (texto) o dict imagen (fallback)
      canonico: metadatos
      chunks: list[str] para map-reduce
      imagenes_selectivas: list[dict] opcionales
    """
    tipo = detectar_tipo_archivo(filename)

    if tipo == "pdf":
        nativo = extraer_pdf_nativo(path)
        if nativo:
            return _pack(nativo)
        ocr = extraer_pdf_ocr(path)
        return _pack(ocr)

    if tipo == "docx":
        return _pack(extraer_docx(path))

    if tipo == "xlsx":
        return _pack(extraer_xlsx(path))

    if tipo == "imagen":
        texto, adv = ocr_archivo_imagen(path)
        if len(texto) >= _TEXTO_MINIMO:
            bloque = BloqueDocumento(id="b1", etiqueta="imagen", texto=texto)
            canon = DocumentoCanonico(
                texto=armar_texto_canonico(
                    meta_lineas=["# meta", "archivo_tipo: imagen", "metodo: rapidocr"],
                    bloques=[bloque],
                ),
                metodo="imagen_ocr",
                paginas=1,
                advertencias=adv,
                bloques=[bloque],
            )
            return _pack(canon)
        # fallback visión
        return {
            "modo": "imagen",
            "contenido": preparar_imagen(path),
            "canonico": {
                "metodo": "imagen_vision",
                "paginas": 1,
                "bloques": 1,
                "advertencias": adv + ["OCR insuficiente; se usa visión"],
                "cobertura_ocr": 0.0,
            },
            "chunks": [],
            "imagenes_selectivas": [],
            "texto_extraido": "",
        }

    raise ValueError(f"Tipo de archivo no soportado: {filename}")


def _pack(canon: DocumentoCanonico) -> dict[str, Any]:
    chunks = chunks_desde_bloques(canon.bloques) if canon.bloques else partir_en_chunks(canon.texto)
    modo: ModoContenido = "mixto" if canon.imagenes_selectivas else "texto"
    return {
        "modo": modo,
        "contenido": canon.texto,
        "texto_extraido": canon.texto,
        "canonico": {
            "metodo": canon.metodo,
            "paginas": canon.paginas,
            "bloques": canon.num_bloques,
            "advertencias": list(canon.advertencias),
            "cobertura_ocr": canon.cobertura_ocr,
        },
        "chunks": chunks,
        "imagenes_selectivas": list(canon.imagenes_selectivas),
    }


# Compat helpers usados por tests/scripts antiguos
def extraer_pdf(path: str) -> str:
    n = extraer_pdf_nativo(path)
    return n.texto if n else ""


def _rasterizar_primera_pagina_pdf(path: str) -> str:
    """Legacy: rasteriza p.1 a archivo temporal PNG."""
    import tempfile

    png = rasterizar_pagina_pdf(path, 0)
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp.write(png)
    tmp.close()
    return tmp.name
