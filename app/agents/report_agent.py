"""Agente reportador: informes por documento y globales + export PDF/DOCX."""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from fpdf import FPDF

from app.core import session_store
from app.core.llm_client import generar_informe_global
from app.models.schemas import DictamenJuridico, DocumentoAnalizado, SesionCompliance


class SesionNoEncontradaError(LookupError):
    pass


class DocumentoNoEncontradoError(LookupError):
    pass


class InformeNoDisponibleError(ValueError):
    pass


def _get_sesion(sesion_id: str) -> SesionCompliance:
    sesion = session_store.get_sesion(sesion_id)
    if sesion is None:
        raise SesionNoEncontradaError(f"Sesión no encontrada: {sesion_id}")
    return sesion


def _get_documento(sesion: SesionCompliance, doc_id: str) -> DocumentoAnalizado:
    for doc in sesion.documentos_analizados:
        if doc.id == doc_id:
            return doc
    raise DocumentoNoEncontradoError(f"Documento no encontrado: {doc_id}")


def informe_documento_markdown(sesion_id: str, doc_id: str) -> str:
    from app.agents.cuestionario import sincronizar_informe_documento

    sesion = _get_sesion(sesion_id)
    doc = _get_documento(sesion, doc_id)
    sincronizar_informe_documento(doc)
    if doc.informe_markdown.strip():
        return doc.informe_markdown
    # Fallback sintético si el LLM no devolvió informe
    obs = "\n".join(
        f"- [{o.severidad}] {o.descripcion}"
        + (f" | Subsanación: {o.subsanacion}" if o.subsanacion else "")
        for o in doc.observaciones
    ) or "- Sin observaciones"
    base = (
        f"# Informe de documento — {doc.tipo.value}\n\n"
        f"**Archivo:** {doc.nombre_archivo}\n\n"
        f"**Cumple:** {'sí' if doc.cumple else 'no'}\n\n"
        f"**Tipo coincide:** {'sí' if getattr(doc, 'tipo_coincide', True) else 'no'}\n\n"
        f"## Resumen\n\n{doc.resumen}\n\n"
        f"## Observaciones\n\n{obs}\n"
    )
    doc.informe_markdown = base
    sincronizar_informe_documento(doc)
    return doc.informe_markdown


def informe_global_markdown(sesion_id: str, *, usar_llm: bool = True) -> str:
    from app.agents.cuestionario import construir_informe_global_estructurado

    sesion = _get_sesion(sesion_id)
    if not sesion.documentos_analizados:
        raise InformeNoDisponibleError(
            "No hay documentos auditados para generar el informe global."
        )
    # Base determinística con todo lo observado + Q&A + recomendaciones
    base = construir_informe_global_estructurado(sesion)
    if not usar_llm:
        return base
    try:
        enriquecido = generar_informe_global(sesion)
        if enriquecido and enriquecido.strip():
            return (
                enriquecido.strip()
                + "\n\n---\n\n## Anexo: detalle estructurado del expediente\n\n"
                + base
            )
    except Exception:
        pass
    return base


def markdown_a_docx_bytes(markdown: str, titulo: str = "Informe") -> bytes:
    doc = Document()
    doc.add_heading(titulo, level=1)
    for linea in markdown.splitlines():
        if linea.startswith("# "):
            doc.add_heading(linea[2:].strip(), level=1)
        elif linea.startswith("## "):
            doc.add_heading(linea[3:].strip(), level=2)
        elif linea.startswith("### "):
            doc.add_heading(linea[4:].strip(), level=3)
        elif linea.startswith("- "):
            doc.add_paragraph(linea[2:], style="List Bullet")
        elif linea.strip():
            doc.add_paragraph(linea)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def markdown_a_pdf_bytes(markdown: str, titulo: str = "Informe") -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 8, _latin1_safe(titulo))
    pdf.ln(4)
    pdf.set_font("Helvetica", size=11)
    for linea in markdown.splitlines():
        texto = _latin1_safe(_strip_md(linea))
        if not texto.strip():
            pdf.ln(4)
            continue
        pdf.set_x(pdf.l_margin)
        if linea.startswith("#"):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, texto.lstrip("#").strip())
            pdf.set_font("Helvetica", size=11)
        else:
            # Evitar líneas enormes sin espacios que rompen el layout
            if len(texto) > 500:
                texto = texto[:500] + "..."
            pdf.multi_cell(0, 6, texto)
        pdf.set_x(pdf.l_margin)
    out = pdf.output()
    if isinstance(out, (bytes, bytearray)):
        return bytes(out)
    return str(out).encode("latin-1", errors="replace")


def _strip_md(linea: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", linea)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s


def _latin1_safe(texto: str) -> str:
    return texto.encode("latin-1", errors="replace").decode("latin-1")


def export_documento(
    sesion_id: str,
    doc_id: str,
    formato: str,
) -> tuple[bytes, str, str]:
    """Devuelve (bytes, media_type, filename)."""
    sesion = _get_sesion(sesion_id)
    doc = _get_documento(sesion, doc_id)
    md = informe_documento_markdown(sesion_id, doc_id)
    titulo = f"Informe {doc.tipo.value}"
    stem = f"informe_{doc.tipo.value.lower()}"
    if formato == "pdf":
        return markdown_a_pdf_bytes(md, titulo), "application/pdf", f"{stem}.pdf"
    if formato == "docx":
        return (
            markdown_a_docx_bytes(md, titulo),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{stem}.docx",
        )
    raise ValueError(f"Formato no soportado: {formato}")


def export_global(sesion_id: str, formato: str, *, usar_llm: bool = True) -> tuple[bytes, str, str]:
    md = informe_global_markdown(sesion_id, usar_llm=usar_llm)
    titulo = "Informe global de auditoría"
    if formato == "pdf":
        return markdown_a_pdf_bytes(md, titulo), "application/pdf", "informe_global.pdf"
    if formato == "docx":
        return (
            markdown_a_docx_bytes(md, titulo),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "informe_global.docx",
        )
    raise ValueError(f"Formato no soportado: {formato}")


def _get_dictamen(sesion: SesionCompliance, dictamen_id: str) -> DictamenJuridico:
    for d in sesion.dictamenes_juridicos or []:
        if d.id == dictamen_id:
            return d
    raise DocumentoNoEncontradoError(f"Dictamen no encontrado: {dictamen_id}")


def informe_juridico_markdown(sesion_id: str, dictamen_id: str) -> str:
    """Dictamen del Jurídico + anexo con hallazgos del Analista (descargable)."""
    from app.agents.cuestionario import construir_informe_global_estructurado

    sesion = _get_sesion(sesion_id)
    dictamen = _get_dictamen(sesion, dictamen_id)
    anexo = construir_informe_global_estructurado(sesion)
    alcance = dictamen.alcance.value if dictamen.alcance else "PARCIAL"
    cabecera = (
        f"# Dictamen jurídico ({alcance})\n\n"
        f"**Expediente:** {sesion.nomenclatura or sesion.id}\n"
        f"**Fecha:** {dictamen.fecha.isoformat()}\n"
        f"**Documentos considerados:** {len(dictamen.documento_ids)}\n"
    )
    if dictamen.slots_pendientes:
        pendientes = ", ".join(
            s.value if hasattr(s, "value") else str(s)
            for s in dictamen.slots_pendientes
        )
        cabecera += f"**Slots pendientes:** {pendientes}\n"
    return (
        f"{cabecera}\n"
        f"{dictamen.markdown.strip()}\n\n"
        f"---\n\n"
        f"## Anexo: hallazgos del Analista y cuestionario\n\n"
        f"{anexo.strip()}\n"
    )


def export_dictamen_juridico(
    sesion_id: str,
    dictamen_id: str,
    formato: str,
) -> tuple[bytes, str, str]:
    """Devuelve (bytes, media_type, filename) del dictamen jurídico exportable."""
    sesion = _get_sesion(sesion_id)
    dictamen = _get_dictamen(sesion, dictamen_id)
    md = informe_juridico_markdown(sesion_id, dictamen_id)
    alcance = dictamen.alcance.value.lower() if dictamen.alcance else "parcial"
    titulo = f"Dictamen jurídico {alcance}"
    stem = f"dictamen_juridico_{alcance}"
    if formato == "pdf":
        return markdown_a_pdf_bytes(md, titulo), "application/pdf", f"{stem}.pdf"
    if formato == "docx":
        return (
            markdown_a_docx_bytes(md, titulo),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{stem}.docx",
        )
    raise ValueError(f"Formato no soportado: {formato}")
